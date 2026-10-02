"""Generate final 5-slide 16:9 presentation for SIG-UOIT incident classifier."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
DELIVERABLES = ROOT / "deliverables"
DELIVERABLES.mkdir(exist_ok=True)
PPTX_PATH = DELIVERABLES / "Presentacion_Clasificacion_Incidentes_UOIT.pptx"
SCREENSHOT_PATH = DELIVERABLES / "screenshot_frontend.png"

# Color Palette: Minimal Tech / Academic Engineering
BG_COLOR = RGBColor(248, 250, 252)       # Light Slate #f8fafc
NAVY_PRIMARY = RGBColor(15, 23, 42)      # Navy Deep #0f172a
NAVY_ACCENT = RGBColor(30, 58, 138)      # Blue 900 #1e3a8a
CYAN_ACCENT = RGBColor(2, 132, 199)      # Cyan / Sky #0284c7
DARK_TEXT = RGBColor(51, 65, 85)         # Slate 700 #334155
MUTED_TEXT = RGBColor(100, 116, 139)     # Slate 500 #64748b
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(255, 255, 255)
BORDER_COLOR = RGBColor(226, 232, 240)   # Slate 200
SUCCESS_BG = RGBColor(240, 253, 244)     # Green tint
SUCCESS_TEXT = RGBColor(21, 128, 61)     # Green text
ACCENT_BG = RGBColor(238, 242, 255)      # Indigo tint


def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_header(slide, title_text, subtitle_text, slide_number):
    # Header shape
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.name = "Arial"
    p.font.color.rgb = NAVY_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.size = Pt(13)
    p2.font.name = "Arial"
    p2.font.color.rgb = CYAN_ACCENT

    # Slide number badge
    num_box = slide.shapes.add_textbox(Inches(12.0), Inches(0.4), Inches(0.8), Inches(0.5))
    ntf = num_box.text_frame
    ntf.margin_left = ntf.margin_top = ntf.margin_right = ntf.margin_bottom = 0
    np_p = ntf.paragraphs[0]
    np_p.text = f"{slide_number} / 5"
    np_p.font.size = Pt(12)
    np_p.font.bold = True
    np_p.font.color.rgb = MUTED_TEXT
    np_p.alignment = PP_ALIGN.RIGHT


def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Problema y Solución
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    add_header(slide1, "Clasificación inteligente de incidentes tecnológicos",
               "Sistema asistido de triaje operativo para la UOIT mediante Machine Learning supervisado", "1")

    # Card 1: Problema
    create_card(slide1, Inches(0.8), Inches(1.5), Inches(5.6), Inches(2.2))
    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(5.2), Inches(1.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "DESAFÍO OPERATIVO: EL PROBLEMA"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf1.add_paragraph()
    p.text = "En la guardia de operaciones e infraestructura (UOIT), el triaje manual de tickets genera cuellos de botella y demoras en la respuesta ante caídas de servicios."
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    p = tf1.add_paragraph()
    p.text = "La clasificación inicial requiere interpretar con rapidez y precisión el título y la descripción preliminar disponibles al momento del registro, antes del diagnóstico técnico."
    p.font.size = Pt(12)
    p.font.color.rgb = MUTED_TEXT

    # Card 2: Solución
    create_card(slide1, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.2))
    tb2 = slide1.shapes.add_textbox(Inches(7.0), Inches(1.65), Inches(5.3), Inches(1.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "SOLUCIÓN PROPUESTA: SISTEMA ASISTIDO"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf2.add_paragraph()
    p.text = "Prototipo funcional desacoplado que recibe el ticket preliminar y sugiere la categoría técnica resolutora para apoyar la toma de decisiones del operador de guardia."
    p.font.size = Pt(13)
    p.font.color.rgb = DARK_TEXT

    p = tf2.add_paragraph()
    p.text = "Entrega un ranking probabilístico ordenado con Top-3 y explicabilidad léxica en tiempo real, operando como herramienta de triaje asistido en CPU (< 20 ms)."
    p.font.size = Pt(12)
    p.font.color.rgb = MUTED_TEXT

    # Card 3: Flujo de Inferencia
    create_card(slide1, Inches(0.8), Inches(3.95), Inches(11.7), Inches(1.3))
    tb_flow = slide1.shapes.add_textbox(Inches(1.0), Inches(4.05), Inches(11.3), Inches(1.1))
    tf_flow = tb_flow.text_frame
    tf_flow.word_wrap = True
    p = tf_flow.paragraphs[0]
    p.text = "FLUJO DE CLASIFICACIÓN ASISTIDA"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY_ACCENT

    p = tf_flow.add_paragraph()
    p.text = "Título + Descripción Inicial   ──▶   Clasificador Supervisado ML   ──▶   Categoría Sugerida (Top-3)   ──▶   Revisión Humana Técnica"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    # Card 4: 5 Categorías Técnicas
    create_card(slide1, Inches(0.8), Inches(5.45), Inches(11.7), Inches(1.5))
    tb_cats = slide1.shapes.add_textbox(Inches(1.0), Inches(5.55), Inches(11.3), Inches(1.3))
    tf_cats = tb_cats.text_frame
    tf_cats.word_wrap = True
    p = tf_cats.paragraphs[0]
    p.text = "CINCO CATEGORÍAS RESOLUTORAS DE INFRAESTRUCTURA UOIT"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    categories = [
        "1. Aplicaciones y servicios (microservicios, APIs, errores 500)",
        "2. Bases de datos (PostgreSQL/Oracle, PgBouncer, réplicas)",
        "3. Redes y conectividad (enlaces WAN, switches VLAN, DNS)",
        "4. Infraestructura y almacenamiento (VMware ESXi, arreglos RAID, SAN)",
        "5. Seguridad y accesos (Directorio Activo, Keycloak SSO, SSL expirados)"
    ]
    p = tf_cats.add_paragraph()
    p.text = "  •  ".join(categories[:3])
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_TEXT

    p = tf_cats.add_paragraph()
    p.text = "  •  ".join(categories[3:])
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_TEXT

    # =========================================================================
    # SLIDE 2: Dataset y Validación
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Dataset y Estrategia de Validación",
               "Generación determinista y prevención estricta de fuga de información (Data Leakage)", "2")

    # 4 Metric Cards for Dataset
    dataset_metrics = [
        ("720", "Incidentes Sintéticos", "100% deterministas (SEED=42)"),
        ("40", "Familias Técnicas", "Disjuntas por event_group_id"),
        ("5", "Categorías Técnicas", "Balance exacto (144 por clase)"),
        ("18", "Variantes por Familia", "Diversidad léxica y sintáctica"),
    ]
    card_w = Inches(2.7)
    for i, (val, title_m, sub_m) in enumerate(dataset_metrics):
        cx = Inches(0.8 + i * 2.98)
        create_card(slide2, cx, Inches(1.5), card_w, Inches(1.4))
        tb = slide2.shapes.add_textbox(cx + Inches(0.15), Inches(1.55), card_w - Inches(0.3), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT
        p = tf.add_paragraph()
        p.text = title_m
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = tf.add_paragraph()
        p.text = sub_m
        p.font.size = Pt(10)
        p.font.color.rgb = MUTED_TEXT

    # Particiones
    create_card(slide2, Inches(0.8), Inches(3.1), Inches(11.7), Inches(1.85))
    tb_part = slide2.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(11.3), Inches(1.65))
    tf_part = tb_part.text_frame
    tf_part.word_wrap = True
    p = tf_part.paragraphs[0]
    p.text = "PARTICIONADO AGRUPADO POR FAMILIAS DISJUNTAS (GROUP SPLIT)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY_ACCENT

    splits = [
        ("ENTRENAMIENTO (Train)", "450 incidentes", "25 familias técnicas (5 por categoría)", "62.5%"),
        ("VALIDACIÓN (Validation)", "90 incidentes", "5 familias técnicas (1 por categoría)", "12.5%"),
        ("PRUEBA FINAL (Test)", "180 incidentes", "10 familias técnicas (2 por categoría)", "25.0%"),
    ]
    for name_s, incs, fams, pct in splits:
        p = tf_part.add_paragraph()
        p.text = f"•  {name_s}: {incs}  |  {fams}  |  Proporción: {pct}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_TEXT

    # Regla Metodológica y Prevención de Fuga
    create_card(slide2, Inches(0.8), Inches(5.15), Inches(11.7), Inches(1.8), bg_color=ACCENT_BG, border_color=CYAN_ACCENT)
    tb_rule = slide2.shapes.add_textbox(Inches(1.0), Inches(5.25), Inches(11.3), Inches(1.6))
    tf_rule = tb_rule.text_frame
    tf_rule.word_wrap = True
    p = tf_rule.paragraphs[0]
    p.text = "PRINCIPIO METODOLÓGICO CLAVE:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY_ACCENT

    p = tf_rule.add_paragraph()
    p.text = "«Las familias permanecen completas en una única partición para evitar fuga de información por memorización de plantillas.»"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = tf_rule.add_paragraph()
    p.text = "• Entrada admitida en inferencia: Exclusivamente título y descripción inicial."
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_TEXT
    p = tf_rule.add_paragraph()
    p.text = "• Prohibición estricta de fuga: Rechazo formal (HTTP 422) de campos como solución, causa raíz o técnico asignado."
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_TEXT

    # =========================================================================
    # SLIDE 3: Comparación Experimental
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Comparación Experimental de Enfoques",
               "Evaluación rigurosa bajo estricto congelamiento metodológico", "3")

    # 3 Column Cards for the 3 Approaches
    col_w = Inches(3.7)
    approaches = [
        ("1. REFERENCIA HEURÍSTICA", "Reglas de palabras clave",
         "Conteo determinista de términos técnicos característicos de la UOIT. Asigna por mayor coincidencia o categoría de respaldo.",
         "Macro-F1 TEST: 0.8793", "Accuracy TEST: 88.89%"),
        ("2. LÍNEA BASE LINEAL", "Dual TF-IDF + LogisticRegression",
         "N-gramas de palabras (1-2) + n-gramas de caracteres (3-5) en FeatureUnion con LogisticRegression (C=1.0).",
         "Macro-F1 TEST: 0.7719", "Accuracy TEST: 78.89% | Top-3: 100%"),
        ("3. MODELO PROPUESTO", "Transformer MiniLM Congelado",
         "Embeddings densos de 384 dimensiones (MiniLM-L12-v2 congelado) + LogisticRegression (C=1.0).",
         "Macro-F1 TEST: 0.6592", "Accuracy TEST: 70.00% | Top-3: 96.1%"),
    ]
    for i, (head, sub_h, desc, f1_str, acc_str) in enumerate(approaches):
        cx = Inches(0.8 + i * 4.0)
        is_selected = (i == 1)
        card_bg = SUCCESS_BG if is_selected else CARD_BG
        b_color = SUCCESS_TEXT if is_selected else BORDER_COLOR
        create_card(slide3, cx, Inches(1.5), col_w, Inches(3.4), bg_color=card_bg, border_color=b_color)
        tb = slide3.shapes.add_textbox(cx + Inches(0.2), Inches(1.6), col_w - Inches(0.4), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = head
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = SUCCESS_TEXT if is_selected else CYAN_ACCENT

        p = tf.add_paragraph()
        p.text = sub_h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_TEXT

        p = tf.add_paragraph()
        p.text = "━━━━━━━━━━━━━━━━━━━━"
        p.font.size = Pt(8)
        p.font.color.rgb = MUTED_TEXT

        p = tf.add_paragraph()
        p.text = f1_str
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY

        p = tf.add_paragraph()
        p.text = acc_str
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = DARK_TEXT

    # Callout: Selección en Validación y Mensaje Principal
    create_card(slide3, Inches(0.8), Inches(5.05), Inches(11.7), Inches(1.9), bg_color=WHITE, border_color=NAVY_ACCENT)
    tb_msg = slide3.shapes.add_textbox(Inches(1.0), Inches(5.15), Inches(11.3), Inches(1.7))
    tf_msg = tb_msg.text_frame
    tf_msg.word_wrap = True

    p = tf_msg.paragraphs[0]
    p.text = "PROTOCOLO DE SELECCIÓN EN VALIDACIÓN Y HALLAZGO TÉCNICO:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = NAVY_ACCENT

    p = tf_msg.add_paragraph()
    p.text = "«Modelo operativo seleccionado previamente en VALIDATION: Línea base — Macro-F1 0.9423 frente a 0.7822 del Transformer.»"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = tf_msg.add_paragraph()
    p.text = "Mensaje Principal: Mayor complejidad no implicó mejor generalización. Los n-gramas de caracteres sub-palabra capturan identificadores especializados (pg_wal, esxi, bgp, raid5) con mayor nitidez que embeddings generales congelados sin ajuste fino."
    p.font.size = Pt(12)
    p.font.color.rgb = DARK_TEXT

    p = tf_msg.add_paragraph()
    p.text = "En Test, la referencia heurística alcanzó el mayor Macro-F1 (0.8793), confirmando el alto valor de reglas léxicas en vocabularios técnicos cerrados. Dicho resultado no fue utilizado para cambiar el modelo operativo."
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED_TEXT

    # =========================================================================
    # SLIDE 4: Sistema Funcional y Demo
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Sistema Funcional y Demostración",
               "Arquitectura de software desacoplada, API REST e interfaz interactiva", "4")

    # Left: Screenshot Real
    if SCREENSHOT_PATH.exists():
        create_card(slide4, Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))
        slide4.shapes.add_picture(str(SCREENSHOT_PATH), Inches(0.85), Inches(1.55), width=Inches(5.9))
    else:
        create_card(slide4, Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.4))

    # Right: Arquitectura y 3 Casos Demostrativos
    create_card(slide4, Inches(7.1), Inches(1.5), Inches(5.4), Inches(2.2))
    tb_arch = slide4.shapes.add_textbox(Inches(7.25), Inches(1.6), Inches(5.1), Inches(2.0))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True

    p = tf_arch.paragraphs[0]
    p.text = "ARQUITECTURA DE SOFTWARE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    p = tf_arch.add_paragraph()
    p.text = "Frontend Web (HTML/JS)  →  FastAPI REST Server"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = tf_arch.add_paragraph()
    p.text = "Validación Pydantic (extra='forbid')  →  Pipeline Scikit-Learn"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    p = tf_arch.add_paragraph()
    p.text = "• Inferencia real en CPU con latencia inferior a 20 ms."
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_TEXT
    p = tf_arch.add_paragraph()
    p.text = "• Desglose probabilístico Top-3 y señales explicativas léxicas."
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_TEXT

    # 3 Demo Cases
    create_card(slide4, Inches(7.1), Inches(3.85), Inches(5.4), Inches(3.05))
    tb_cases = slide4.shapes.add_textbox(Inches(7.25), Inches(3.95), Inches(5.1), Inches(2.85))
    tf_cases = tb_cases.text_frame
    tf_cases.word_wrap = True

    p = tf_cases.paragraphs[0]
    p.text = "TRES CASOS DEMOSTRATIVOS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    cases_info = [
        ("A. Caso Válido (Bases de Datos):",
         "Saturación de PgBouncer en PostgreSQL  →  Inferencia con puntaje > 80% en 'Bases de datos' y señales claras."),
        ("B. Caso Ambiguo (Competencia de Clases):",
         "Error 500 y timeout en reportes  →  Puntajes repartidos entre 'Aplicaciones' y 'Bases de datos' en Top-3."),
        ("C. Entrada Inválida (Violación de Contrato):",
         "Texto telegráfico ('Caída' / 'Error')  →  Interceptado formalmente con código HTTP 422.")
    ]
    for c_title, c_desc in cases_info:
        p = tf_cases.add_paragraph()
        p.text = c_title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p = tf_cases.add_paragraph()
        p.text = c_desc
        p.font.size = Pt(10)
        p.font.color.rgb = DARK_TEXT

    # =========================================================================
    # SLIDE 5: Conclusiones y Limitaciones
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Conclusiones, Limitaciones y Repositorio",
               "Síntesis académica y consideraciones para la adopción institucional", "5")

    # Left: Conclusiones
    create_card(slide5, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.3))
    tb_c = slide5.shapes.add_textbox(Inches(1.0), Inches(1.65), Inches(5.2), Inches(4.0))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "CONCLUSIONES TÉCNICAS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY_ACCENT

    conclusions = [
        "1. Flujo Completo Implementado: Se diseñó, validó y desplegó el pipeline supervisado completo bajo estricto congelamiento.",
        "2. Superioridad de la Línea Base: La extracción sub-palabra (char_wb) superó a los embeddings congelados del Transformer en Validación (0.9423 vs 0.7822) y Test.",
        "3. Efectividad Heurística: Las reglas simples obtuvieron el mayor Macro-F1 en Test (0.8793), demostrando que mayor complejidad no asegura mejor generalización.",
        "4. Exactitud Top-3 del 100%: En el 100% de los incidentes de test, el área resolutora correcta estuvo dentro de las 3 opciones sugeridas al operador."
    ]
    for item in conclusions:
        p = tf_c.add_paragraph()
        p.text = item
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_TEXT

    # Right: Limitaciones
    create_card(slide5, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.3))
    tb_l = slide5.shapes.add_textbox(Inches(7.0), Inches(1.65), Inches(5.3), Inches(4.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "LIMITACIONES METODOLÓGICAS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    limitations = [
        "1. Dataset Sintético: 720 incidentes en 40 familias. Presenta mayor regularidad y separabilidad léxica que tickets humanos reales.",
        "2. Taxonomía Cerrada: Clasificación mono-etiqueta en 5 categorías primarias. No modela fallos compuestos simultáneos.",
        "3. Transformer Congelado: Evaluado sin fine-tuning supervisado por restricciones de cómputo en CPU.",
        "4. Puntajes No Calibrados: Las salidas son puntajes relativos del modelo y no probabilidades calibradas de acierto absoluto.",
        "5. Validación en Piloto Requerida: Los resultados no demuestran eficacia en producción sin un piloto previo con incidentes reales anonimizados."
    ]
    for item in limitations:
        p = tf_l.add_paragraph()
        p.text = item
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_TEXT

    # Bottom: Cierre y Repositorio
    create_card(slide5, Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.0), bg_color=NAVY_PRIMARY, border_color=None)
    tb_bot = slide5.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.9))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True

    p = tf_bot.paragraphs[0]
    p.text = "«El sistema asiste al operador; no reemplaza su criterio técnico.»"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    p = tf_bot.add_paragraph()
    p.text = "Repositorio Oficial: https://github.com/juan-serrudo/sig-incidentes   |   34 pruebas pytest aprobadas"
    p.font.size = Pt(11)
    p.font.color.rgb = CYAN_ACCENT
    p.alignment = PP_ALIGN.CENTER

    prs.save(str(PPTX_PATH))
    print(f"Presentación PPTX generada con éxito en: {PPTX_PATH}")


if __name__ == "__main__":
    build_presentation()
